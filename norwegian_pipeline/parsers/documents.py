"""Parsers for tender attachment formats."""
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from norwegian_pipeline.matching import match_text
from norwegian_pipeline.models import PackRow
from norwegian_pipeline.normalisation import parse_norwegian_number

logger = logging.getLogger(__name__)

# Column header aliases (Norwegian + English) mapped to internal fields
HEADER_MAP: dict[str, str] = {
    "atc": "atc_code",
    "atc-kode": "atc_code",
    "atc kode": "atc_code",
    "atc_code": "atc_code",
    "varenummer": "item_number",
    "item number": "item_number",
    "itemnummer": "item_number",
    "produktnavn": "product_name",
    "product name": "product_name",
    "preparat": "product_name",
    "styrke": "strength",
    "strength": "strength",
    "pakning": "packsize",
    "pakningsstørrelse": "packsize",
    "pack size": "packsize",
    "leverandør": "supplier",
    "supplier": "supplier",
    "makspris": "max_price",
    "max price": "max_price",
    "pris": "max_price",
    "forbruk": "packs_sold_last_12m",
    "forbruk siste 12 mnd": "packs_sold_last_12m",
    "volum": "packs_sold_last_12m",
    "virkestoff": "product_name",
}


def _normalise_header(cell: str) -> str:
    key = re.sub(r"\s+", " ", (cell or "").strip().lower())
    return HEADER_MAP.get(key, "")


def _document_detection_method(match_method: str) -> str:
    """Map internal match methods to CSV detectionMethod values."""
    mapping = {
        "atc_code": "atc_code_in_document",
        "name_in_text": "name_in_document",
        "brand_name": "brand_in_document",
        "lis_ref_in_title": "lis_ref_in_document",
    }
    return mapping.get(match_method, f"{match_method}_in_document")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_xlsx(path: Path, *, notice_id: str = "", source_document: str = "") -> list[PackRow]:
    rows: list[PackRow] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("Could not parse xlsx %s: %s", path, exc)
        return rows

    for sheet in wb.worksheets:
        header_row_idx = None
        headers: dict[int, str] = {}
        for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            mapping: dict[int, str] = {}
            for col_idx, cell in enumerate(row):
                field = _normalise_header(_cell_str(cell))
                if field:
                    mapping[col_idx] = field
            if len(mapping) >= 2:
                header_row_idx = i
                headers = mapping
                break

        if header_row_idx is None:
            continue

        for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
            data: dict[str, str] = {}
            row_text_parts: list[str] = []
            for col_idx, cell in enumerate(row):
                text = _cell_str(cell)
                if text:
                    row_text_parts.append(text)
                field = headers.get(col_idx)
                if field and text:
                    data[field] = text

            if not row_text_parts:
                continue

            row_text = " ".join(row_text_parts)
            matches = match_text(row_text)
            if not matches:
                continue

            for hit in matches:
                rows.append(
                    PackRow(
                        notice_id=notice_id,
                        product_molecule=hit.molecule,
                        molecule_detected="yes" if hit.detected else "no",
                        molecule_variant=hit.variant,
                        detection_method=_document_detection_method(hit.method),
                        atc_code=data.get("atc_code", hit.atc_code),
                        item_number=data.get("item_number", ""),
                        product_name=data.get("product_name", ""),
                        strength=data.get("strength", ""),
                        packsize=data.get("packsize", ""),
                        supplier=data.get("supplier", ""),
                        max_price=_cell_str(parse_norwegian_number(data.get("max_price", ""))),
                        packs_sold_last_12m=data.get("packs_sold_last_12m", ""),
                        source_document=source_document or path.name,
                    )
                )
    wb.close()
    return rows


def parse_pdf(path: Path, *, notice_id: str = "", source_document: str = "") -> list[PackRow]:
    rows: list[PackRow] = []
    try:
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        logger.warning("Could not parse pdf %s: %s", path, exc)
        return rows

    if not text.strip():
        return rows

    for hit in match_text(text):
        rows.append(
            PackRow(
                notice_id=notice_id,
                product_molecule=hit.molecule,
                molecule_detected="yes",
                molecule_variant=hit.variant,
                detection_method=_document_detection_method(hit.method),
                atc_code=hit.atc_code,
                source_document=source_document or path.name,
            )
        )
    return rows


def parse_document(path: Path, **kwargs: Any) -> list[PackRow]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx(path, **kwargs)
    if suffix == ".pdf":
        return parse_pdf(path, **kwargs)
    if suffix == ".html":
        text = path.read_text(encoding="utf-8", errors="replace")
        return _rows_from_text(text, **kwargs)
    return []


def _rows_from_text(text: str, **kwargs: Any) -> list[PackRow]:
    notice_id = kwargs.get("notice_id", "")
    source_document = kwargs.get("source_document", "")
    rows: list[PackRow] = []
    for hit in match_text(text):
        rows.append(
            PackRow(
                notice_id=notice_id,
                product_molecule=hit.molecule,
                molecule_detected="yes",
                molecule_variant=hit.variant,
                detection_method=_document_detection_method(hit.method),
                atc_code=hit.atc_code,
                source_document=source_document,
            )
        )
    return rows
